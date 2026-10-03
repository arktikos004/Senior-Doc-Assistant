"""對帳彙整匯出(export_ledger)的單元測試。"""
import openpyxl

from src.export_ledger import COLUMNS, build_rows, export_ledger, totals_by_currency


def _rec(fname, when, date, vendor, amount, action="archive", conf=0.9,
         doc_type="發票", currency="NTD"):
    return {
        "時間": when,
        "原始檔案": fname,
        "動作": action,
        "目標路徑": f"data/archive/{doc_type}/2026-07/{fname}",
        "AI辨識結果": {
            "doc_type": doc_type, "date": date, "vendor": vendor, "amount": amount,
            "currency": currency, "invoice_number": "AB12345678", "confidence": conf,
            "notes": "",
        },
    }


def test_build_rows_dedupes_keeping_latest():
    records = [
        _rec("a.pdf", "2026-08-13T10:00:00", "2026-07-01", "舊", 100),
        _rec("a.pdf", "2026-08-13T11:00:00", "2026-07-01", "新", 200),  # 較新→勝出
        _rec("b.pdf", "2026-08-13T10:00:00", "2026-06-01", "B", 50),
    ]
    rows = build_rows(records)
    assert len(rows) == 2  # a.pdf 去重
    a = next(r for r in rows if r["原始檔名"] == "a.pdf")
    assert a["商家"] == "新" and a["金額"] == 200  # 取較新那筆


def test_build_rows_filters_by_doc_type():
    records = [
        _rec("inv.pdf", "t", "2026-07-01", "V", 100, doc_type="發票"),
        _rec("rec.pdf", "t", "2026-07-01", "V", 100, doc_type="收據"),
    ]
    rows = build_rows(records, doc_type="發票")
    assert [r["原始檔名"] for r in rows] == ["inv.pdf"]


def test_totals_split_by_currency():
    records = [
        _rec("a.pdf", "t", "2026-07-01", "A", 1000, currency="NTD"),
        _rec("b.pdf", "t", "2026-07-02", "B", 20, currency="USD"),
        _rec("c.pdf", "t", "2026-07-03", "C", 500, currency="NTD"),
    ]
    totals = totals_by_currency(build_rows(records))
    assert totals["NTD"] == 1500 and totals["USD"] == 20  # 幣別不混算


def test_export_ledger_writes_xlsx_with_currency_subtotals(tmp_path):
    records = [
        _rec("a.pdf", "t", "2026-07-01", "商家A", 100, currency="NTD"),
        _rec("b.pdf", "t", "2026-07-02", "商家B", 20, currency="USD"),
    ]
    out = tmp_path / "對帳.xlsx"
    path, n, totals = export_ledger(records, out)
    assert path.exists() and n == 2
    assert totals == {"NTD": 100, "USD": 20}

    ws = openpyxl.load_workbook(out).active
    assert [c.value for c in ws[1]] == COLUMNS  # 表頭
    rows = list(ws.iter_rows(values_only=True))
    vendor_i, amt_i = COLUMNS.index("商家"), COLUMNS.index("金額")
    labels = {r[vendor_i]: r[amt_i] for r in rows if r[vendor_i] and str(r[vendor_i]).startswith("合計")}
    assert labels == {"合計 (NTD)": 100, "合計 (USD)": 20}
